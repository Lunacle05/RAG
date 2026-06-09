from __future__ import annotations

import os
from dataclasses import dataclass
from typing import Dict, Iterable, List, Tuple

from docx import Document
from pypdf import PdfReader
from openpyxl import load_workbook
from PIL import Image

try:  # 可选依赖：PDF 表格与文字分离
    import pdfplumber
    _PDFPLUMBER_AVAILABLE = True
except Exception:
    _PDFPLUMBER_AVAILABLE = False

try:  # 可选依赖：从 PDF 中提取内嵌图片
    import fitz as pymupdf
    _PYMUPDF_AVAILABLE = True
except Exception:
    _PYMUPDF_AVAILABLE = False

try:  # 可选依赖：图片内文字识别，使返回的 content 为图片中的实际文字
    import pytesseract
    _OCR_AVAILABLE = True
except Exception:
    _OCR_AVAILABLE = False


def _ocr_image(path: str) -> str:
    """从图片路径识别文字（中英），无依赖或失败时返回空字符串。"""
    if not _OCR_AVAILABLE:
        return ""
    try:
        img = Image.open(path)
        # 中英混合场景，需系统安装 tesseract 及 chi_sim 语言包
        text = pytesseract.image_to_string(img, lang="chi_sim+eng")
        return (text or "").strip()
    except Exception:
        return ""


@dataclass(frozen=True)
class LoadedDoc:
    text: str
    metadata: Dict


SUPPORTED_EXTS = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
    ".xlsx",
    ".xls",
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
}


# 入库时生成的 PDF 页图缓存目录，不当作用户资料参与扫描
_RAG_PDF_IMAGES_DIR = ".rag_pdf_images"


def iter_files(input_path: str) -> Iterable[str]:
    if os.path.isfile(input_path):
        yield input_path
        return
    for root, _, files in os.walk(input_path):
        # 跳过“PDF 转页图”生成的缓存目录，只统计用户真正放入的文件
        if _RAG_PDF_IMAGES_DIR in root:
            continue
        for f in files:
            p = os.path.join(root, f)
            ext = os.path.splitext(p)[1].lower()
            if ext in SUPPORTED_EXTS:
                yield p


def _read_text_file(path: str) -> str:
    # Common encodings for CN/EN Windows environments
    for enc in ("utf-8", "utf-8-sig", "gb18030", "gbk"):
        try:
            with open(path, "r", encoding=enc) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def _extract_nearby_text_for_image(page, img_rect, max_up: int = 2, max_down: int = 2, max_chars: int = 800):
    """
    提取图片附近的文本块：
    - 上方取最近的若干块
    - 下方取最近的若干块
    返回 (upper_text, lower_text)
    """
    try:
        blocks = page.get_text("blocks") or []
    except Exception:
        return "", ""

    ix0, iy0, ix1, iy1 = img_rect.x0, img_rect.y0, img_rect.x1, img_rect.y1

    upper = []
    lower = []

    for block in blocks:
        if len(block) < 5:
            continue

        x0, y0, x1, y1, text = block[:5]
        text = (text or "").strip()
        if not text:
            continue

        overlap_x = max(0, min(ix1, x1) - max(ix0, x0))
        horiz_overlap = overlap_x > 0

        block_center_x = (x0 + x1) / 2
        image_center_x = (ix0 + ix1) / 2
        center_dx = abs(block_center_x - image_center_x)

        # 图片上方
        if y1 <= iy0:
            vertical_distance = iy0 - y1
            if horiz_overlap or vertical_distance <= 160:
                score = (vertical_distance, center_dx, y0)
                upper.append((score, text))

        # 图片下方
        elif y0 >= iy1:
            vertical_distance = y0 - iy1
            if horiz_overlap or vertical_distance <= 160:
                score = (vertical_distance, center_dx, y0)
                lower.append((score, text))

    upper.sort(key=lambda x: x[0])
    lower.sort(key=lambda x: x[0])

    upper_picked = []
    lower_picked = []

    upper_len = 0
    for _, text in upper[:max_up]:
        if text not in upper_picked:
            upper_picked.append(text)
            upper_len += len(text)

    lower_len = 0
    for _, text in lower[:max_down]:
        if text not in lower_picked:
            lower_picked.append(text)
            lower_len += len(text)

    upper_text = "\n".join(upper_picked).strip()
    lower_text = "\n".join(lower_picked).strip()

    # 简单限制总长度
    total = (upper_text + "\n" + lower_text).strip()
    if len(total) > max_chars:
        remain = max_chars
        if upper_text:
            upper_text = upper_text[: remain // 2].strip()
        if lower_text:
            lower_text = lower_text[: remain // 2].strip()

    return upper_text, lower_text


def load_one(path: str) -> List[LoadedDoc]:
    ext = os.path.splitext(path)[1].lower()
    source = os.path.basename(path)

    if ext == ".pdf":
        docs: List[LoadedDoc] = []
        abs_path = os.path.abspath(path)

        # 1) 文字与表格分离：优先用 pdfplumber 分别提取文字和表格
        if _PDFPLUMBER_AVAILABLE:
            try:
                with pdfplumber.open(path) as pdf:
                    for i, page in enumerate(pdf.pages):
                        # 纯文字
                        page_text = (page.extract_text() or "").strip()
                        if page_text:
                            docs.append(
                                LoadedDoc(
                                    text=page_text,
                                    metadata={
                                        "source": source,
                                        "source_path": abs_path,
                                        "page": i + 1,
                                        "type": "pdf",
                                    },
                                )
                            )
                        # 表格：像 Excel 一样，按“表头 + 行”展开，便于 LLM 理解列含义；
                        # 同时在文本前附带一两行页面上下文，帮助理解表格所在语境。
                        tables = page.extract_tables() or []
                        for t_idx, table in enumerate(tables):
                            if not table:
                                continue
                            rows_list = [
                                [str(cell or "").strip() for cell in row]
                                for row in table
                                if any(cell for cell in row)
                            ]
                            # 只有表头没有数据行时，不单独为表头建向量
                            if len(rows_list) <= 1:
                                continue
                            # 取页面文本的前一两行作为上下文
                            context_lines = []
                            if page_text:
                                for line in page_text.split("\n"):
                                    line = line.strip()
                                    if line:
                                        context_lines.append(line)
                                    if len(context_lines) >= 2:
                                        break
                            context_prefix = "\n".join(context_lines) if context_lines else ""

                            header_line = " | ".join(rows_list[0])
                            # 从第 2 行开始，每行一条：表头 + 当前行
                            for row_idx, cells in enumerate(rows_list[1:], start=2):
                                row_line = " | ".join(cells)
                                text_row_core = header_line + "\n" + row_line
                                text_row = (
                                    context_prefix + "\n" + text_row_core
                                    if context_prefix
                                    else text_row_core
                                )
                                docs.append(
                                    LoadedDoc(
                                        text=text_row,
                                        metadata={
                                            "source": source,
                                            "source_path": abs_path,
                                            "page": i + 1,
                                            "type": "pdf_table",
                                            "table_index": t_idx,
                                            "row_index": row_idx,
                                        },
                                    )
                                )
            except Exception:
                pass

        # 若无 pdfplumber 或提取失败，回退到 pypdf 仅提取文字
        if not docs:
            reader = PdfReader(path)
            for i, page in enumerate(reader.pages):
                text = (page.extract_text() or "").strip()
                if text:
                    docs.append(
                        LoadedDoc(
                            text=text,
                            metadata={
                                "source": source,
                                "source_path": abs_path,
                                "page": i + 1,
                                "type": "pdf",
                            },
                        )
                    )

        # 2) 内嵌图片：用 PyMuPDF 提取，按“照片”方式处理（VL 向量 + OCR 文字）
        if _PYMUPDF_AVAILABLE:
            try:
                img_dir = os.path.join(os.path.dirname(path), _RAG_PDF_IMAGES_DIR)
                os.makedirs(img_dir, exist_ok=True)
                doc_fitz = pymupdf.open(path)
                for page_no in range(len(doc_fitz)):
                    page = doc_fitz[page_no]
                    for img_idx, img_info in enumerate(page.get_images(full=True)):
                        xref = img_info[0]
                        try:
                            base_img = doc_fitz.extract_image(xref)
                            img_bytes = base_img["image"]
                            ext_img = base_img.get("ext", "png")
                            img_name = f"{os.path.splitext(source)[0]}_p{page_no + 1}_img{img_idx}.{ext_img}"
                            img_path = os.path.join(img_dir, img_name)
                            with open(img_path, "wb") as f:
                                f.write(img_bytes)
                            abs_img = os.path.abspath(img_path)

                            ocr_text = _ocr_image(abs_img)

                            upper_text = ""
                            lower_text = ""
                            try:
                                rects = page.get_image_rects(xref)
                                if rects:
                                    upper_text, lower_text = _extract_nearby_text_for_image(page, rects[0])
                            except Exception:
                                upper_text, lower_text = "", ""

                            # 找不到上下附近文本时，回退到整页前几行作为上方文本
                            if not upper_text and not lower_text:
                                page_text_for_image = (page.get_text("text") or "").strip()
                                top_lines = [line.strip() for line in page_text_for_image.splitlines() if line.strip()]
                                if top_lines:
                                    upper_text = "\n".join(top_lines[:6])

                            parts = []

                            if upper_text:
                                parts.append(upper_text)

                            if ocr_text:
                                parts.append(ocr_text[:1500])

                            if lower_text:
                                parts.append(lower_text)

                            final_text = "\n\n".join(parts).strip()


                            docs.append(
                                LoadedDoc(
                                    text=final_text,
                                    metadata={
                                        "source": source,
                                        "source_path": abs_path,
                                        "page": page_no + 1,
                                        "type": "image",
                                        "image_path": abs_img,
                                        "image_index": img_idx,
                                    },
                                )
                            )
                        except Exception:
                            continue
                doc_fitz.close()
            except Exception:
                pass

        return docs

    if ext == ".docx":
        doc = Document(path)
        abs_path = os.path.abspath(path)

        def _is_heading_style(style_name: str) -> bool:
            if not style_name:
                return False
            return "Heading" in style_name or style_name.startswith("标题")

        sections: List[Tuple[str, str]] = []  # (text, type: "docx")
        current: List[str] = []
        for p in doc.paragraphs:
            t = (p.text or "").strip()
            style_name = getattr(p.style, "name", "") or ""
            if _is_heading_style(style_name):
                if current:
                    sections.append(("\n".join(current), "docx"))
                    current = []
            if t:
                current.append(t)
        if current:
            sections.append(("\n".join(current), "docx"))

        docs_docx: List[LoadedDoc] = []
        for i, (sec, sec_type) in enumerate(sections):
            if not sec.strip():
                continue
            docs_docx.append(
                LoadedDoc(
                    text=sec.strip(),
                    metadata={
                        "source": source,
                        "source_path": abs_path,
                        "type": sec_type,
                        "section_index": i,
                    },
                )
            )

        # 记录每个表格前最近一段文本，作为“表格名/上下文”
        table_context_by_element: Dict = {}
        last_text_para = ""
        for child in doc.element.body.iterchildren():
            tag = str(child.tag)
            if tag.endswith("}p"):
                txt = "".join(child.itertext()).strip()
                if txt:
                    last_text_para = txt
            elif tag.endswith("}tbl"):
                table_context_by_element[child] = last_text_para

        # Word 表格：每条数据一块，携带表格上下文/名称 + 表头 + 当前行
        for table_idx, tbl in enumerate(doc.tables):
            rows_list: List[List[str]] = []
            for row in tbl.rows:
                cells = [str(c.text or "").strip() for c in row.cells]
                if any(cells):
                    rows_list.append(cells)

            # 只有表头没有数据行时，不单独建向量
            if len(rows_list) <= 1:
                continue

            context_text = str(table_context_by_element.get(tbl._element, "") or "").strip()
            table_name = context_text or f"表格{table_idx + 1}"
            prefix = f"文档: {source} | 表格: {table_name}"
            header_line = " | ".join(rows_list[0])

            for row_idx, cells in enumerate(rows_list[1:], start=2):
                row_line = " | ".join(cells)
                docs_docx.append(
                    LoadedDoc(
                        text=prefix + "\n" + header_line + "\n" + row_line,
                        metadata={
                            "source": source,
                            "source_path": abs_path,
                            "type": "docx_table",
                            "table_index": table_idx,
                            "row_index": row_idx,
                            "table_name": table_name,
                        },
                    )
                )

        return docs_docx


    if ext in {".txt", ".md"}:
        text = _read_text_file(path).strip()
        return [
            LoadedDoc(
                text=text,
                metadata={
                    "source": source,
                    "source_path": os.path.abspath(path),
                    "type": ext.lstrip("."),
                },
            )
        ]

    if ext in {".xlsx", ".xls"}:
        wb = load_workbook(path, data_only=True)
        docs_excel: List[LoadedDoc] = []
        abs_path = os.path.abspath(path)
        for sheet in wb.worksheets:
            rows_list: List[List[str]] = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c not in (None, "")]
                if cells:
                    rows_list.append(cells)
            # 只有表头没有数据行时，不单独为表头建向量
            if len(rows_list) <= 1:
                continue
            header_line = " | ".join(rows_list[0])
            prefix = f"表格文件: {source} | 工作表: {sheet.title}"
            # 从第 2 行开始，每行一条：文件/表名 + 表头 + 当前行
            for row_idx, cells in enumerate(rows_list[1:], start=2):
                row_line = " | ".join(cells)
                text = prefix + "\n" + header_line + "\n" + row_line
                docs_excel.append(
                    LoadedDoc(
                        text=text,
                        metadata={
                            "source": source,
                            "source_path": abs_path,
                            "type": "excel",
                            "sheet": sheet.title,
                            "row_index": row_idx,
                        },
                    )
                )
        return docs_excel

    if ext in {".png", ".jpg", ".jpeg", ".webp"}:
        # 用 OCR 提取图片内文字作为 content，便于检索返回“图片里的具体内容”；向量仍用 Qwen3-VL 图片通路。
        abs_path = os.path.abspath(path)
        ocr_text = _ocr_image(abs_path)
        return [
            LoadedDoc(
                text=ocr_text,
                metadata={
                    "source": source,
                    "source_path": abs_path,
                    "type": "image",
                    "image_path": abs_path,
                },
            )
        ]

    raise ValueError(f"Unsupported file type: {ext}")


def load_many(input_path: str) -> List[LoadedDoc]:
    docs: List[LoadedDoc] = []
    for fp in iter_files(input_path):
        docs.extend(load_one(fp))
    return docs